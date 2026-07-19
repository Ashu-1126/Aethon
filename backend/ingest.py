"""
AETHON — Document Loader & Chunker
Stage 1 of the ingestion pipeline.

Supports: PDF (text + OCR fallback), DOCX, TXT, XLSX/CSV
Each chunk carries metadata: doc_name, page, chunk_index, doc_type
"""
from __future__ import annotations

import re
import io
from pathlib import Path
from typing import Generator

from config import CHUNK_SIZE, CHUNK_OVERLAP


# ── Type alias ──────────────────────────────────────────────────────────────
Chunk = dict  # {text, doc_name, page, chunk_index, doc_type}


# ══════════════════════════════════════════════════════════════════════════
# PARSERS
# ══════════════════════════════════════════════════════════════════════════

def _parse_pdf(path: Path) -> list[tuple[int, str]]:
    """Return list of (page_number, text) tuples.  Falls back to OCR if a
    page has < 30 chars of extractable text (i.e., it's a scanned image)."""
    import fitz  # PyMuPDF

    pages: list[tuple[int, str]] = []
    doc = fitz.open(str(path))
    for i, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        if len(text) < 30:
            # OCR fallback
            text = _ocr_page(page)
        if text:
            pages.append((i, text))
    doc.close()
    return pages


def _ocr_page(page) -> str:  # fitz.Page
    """Render PDF page to image and run Tesseract OCR."""
    try:
        import pytesseract
        from PIL import Image

        mat = page.get_pixmap(dpi=200)
        img_bytes = mat.tobytes("png")
        img = Image.open(io.BytesIO(img_bytes))
        return pytesseract.image_to_string(img).strip()
    except Exception:
        return ""


def _parse_docx(path: Path) -> list[tuple[int, str]]:
    """DOCX: treat every 10 paragraphs as a logical 'page'."""
    from docx import Document

    doc = Document(str(path))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    pages: list[tuple[int, str]] = []
    page_size = 10
    for i in range(0, len(paras), page_size):
        page_num = i // page_size + 1
        text = "\n".join(paras[i : i + page_size])
        pages.append((page_num, text))
    return pages


def _parse_txt(path: Path) -> list[tuple[int, str]]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    lines = text.splitlines()
    pages = []
    size = 60  # lines per virtual page
    for i in range(0, len(lines), size):
        pages.append((i // size + 1, "\n".join(lines[i : i + size])))
    return pages


def _parse_csv(path: Path) -> list[tuple[int, str]]:
    import csv

    rows = []
    with open(path, newline="", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append(", ".join(row))
    pages = []
    size = 30
    for i in range(0, len(rows), size):
        pages.append((i // size + 1, "\n".join(rows[i : i + size])))
    return pages


def _parse_xlsx(path: Path) -> list[tuple[int, str]]:
    """Convert each sheet to CSV-like text."""
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        pages = []
        for sheet_idx, sheet in enumerate(wb.worksheets, start=1):
            rows = []
            for row in sheet.iter_rows(values_only=True):
                row_text = ", ".join(str(c) if c is not None else "" for c in row)
                if row_text.strip(", "):
                    rows.append(row_text)
            if rows:
                pages.append((sheet_idx, "\n".join(rows)))
        wb.close()
        return pages
    except Exception as e:
        print(f"[Ingest Error] Failed to parse XLSX at {path.name}: {e}")
        return []


def _parse_html(path: Path) -> list[tuple[int, str]]:
    from html.parser import HTMLParser

    class HTMLTextExtractor(HTMLParser):
        def __init__(self):
            super().__init__()
            self.text_parts = []
            self.ignore_tags = {"script", "style", "head", "meta", "title", "link"}
            # Stack of currently-open tags. Data is ignored only while an
            # ancestor tag is in `ignore_tags` (e.g. text inside <script>).
            self.tag_stack = []

        def handle_starttag(self, tag, attrs):
            self.tag_stack.append(tag)

        def handle_endtag(self, tag):
            # Pop the matching open tag (handles malformed/mismatched HTML).
            if tag in self.tag_stack:
                idx = len(self.tag_stack) - 1 - self.tag_stack[::-1].index(tag)
                del self.tag_stack[idx:]

        def handle_data(self, data):
            if not any(t in self.ignore_tags for t in self.tag_stack):
                text = data.strip()
                if text:
                    self.text_parts.append(text)

        def get_text(self):
            return "\n".join(self.text_parts)

    try:
        raw_html = path.read_text(encoding="utf-8", errors="ignore")
        extractor = HTMLTextExtractor()
        extractor.feed(raw_html)
        text = extractor.get_text()
        
        # Split html text into logical page-sized blocks of 40 lines each
        lines = text.splitlines()
        pages = []
        size = 40
        for i in range(0, len(lines), size):
            pages.append((i // size + 1, "\n".join(lines[i : i + size])))
        return pages
    except Exception:
        return []


def _parse_image(path: Path) -> list[tuple[int, str]]:
    """Render image and run Tesseract OCR."""
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(path)
        text = pytesseract.image_to_string(img).strip()
        if text:
            return [(1, text)]  # virtual page 1
    except Exception as e:
        print(f"[OCR Warning] Image parsing failed for {path.name}: {e}")
    return []


PARSERS = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".txt": _parse_txt,
    ".csv": _parse_csv,
    ".xlsx": _parse_xlsx,
    ".html": _parse_html,
    ".htm": _parse_html,
    ".png": _parse_image,
    ".jpg": _parse_image,
    ".jpeg": _parse_image,
}


def _infer_doc_type(name: str) -> str:
    n = name.lower()
    if any(x in n for x in ["regulation", "act", "rule", "oisd", "dgms", "peso"]):
        return "regulation"
    if any(x in n for x in ["sop", "procedure", "protocol", "wip"]):
        return "procedure"
    if any(x in n for x in ["manual", "oem", "datasheet"]):
        return "manual"
    if any(x in n for x in ["incident", "accident", "near", "hazard"]):
        return "incident"
    if any(x in n for x in ["pid", "drawing", "layout", "p&id"]):
        return "drawing"
    if any(x in n for x in ["work", "order", "wo", "log"]):
        return "document"
    return "document"


# ══════════════════════════════════════════════════════════════════════════
# CHUNKER  (approx-token, overlap)
# ══════════════════════════════════════════════════════════════════════════

def _approx_tokens(text: str) -> int:
    """~4 chars per token for English industrial text."""
    return len(text) // 4


def _split_page(text: str, target: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split page text into overlapping chunks of ~target tokens."""
    sentences = re.split(r"(?<=[.!?])\s+", text)
    chunks: list[str] = []
    current: list[str] = []
    current_tokens = 0

    for sent in sentences:
        t = _approx_tokens(sent)
        if current_tokens + t > target and current:
            chunks.append(" ".join(current))
            # keep overlap sentences
            overlap_sents: list[str] = []
            ov = 0
            for s in reversed(current):
                ov += _approx_tokens(s)
                overlap_sents.insert(0, s)
                if ov >= overlap:
                    break
            current = overlap_sents
            current_tokens = sum(_approx_tokens(s) for s in current)
        current.append(sent)
        current_tokens += t

    if current:
        chunks.append(" ".join(current))
    return [c.strip() for c in chunks if c.strip()]


# ══════════════════════════════════════════════════════════════════════════
# PUBLIC API
# ══════════════════════════════════════════════════════════════════════════

def load_and_chunk(path: Path) -> list[Chunk]:
    """
    Parse a document and return a flat list of Chunks.
    Each chunk is a dict ready for embedding + storage.
    """
    suffix = path.suffix.lower()
    parser = PARSERS.get(suffix)
    if parser is None:
        raise ValueError(f"Unsupported file type: {suffix}")

    pages = parser(path)
    doc_name = path.name
    doc_type = _infer_doc_type(doc_name)

    chunks: list[Chunk] = []
    chunk_idx = 0
    for page_num, page_text in pages:
        for sub_text in _split_page(page_text):
            chunks.append(
                {
                    "text": sub_text,
                    "doc_name": doc_name,
                    "page": page_num,
                    "chunk_index": chunk_idx,
                    "doc_type": doc_type,
                }
            )
            chunk_idx += 1

    return chunks


# ══════════════════════════════════════════════════════════════════════════
# CLI helper — python ingest.py path/to/doc.pdf
# ══════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    import sys, json

    p = Path(sys.argv[1])
    result = load_and_chunk(p)
    print(f"Loaded {len(result)} chunks from {p.name}")
    if result:
        print(json.dumps(result[0], indent=2))
